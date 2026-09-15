"""Export an immutable work product; keep provenance in a separate appendix."""
import json
import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn


def build(payload, destination):
    doc = Document()
    # The runtime's default Word template can carry a blue Title border.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.85)
    section.left_margin = section.right_margin = Inches(0.95)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = doc.styles[name]
        style.font.name = '宋体'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体' if name == 'Normal' else '黑体')
    normal = doc.styles['Normal']
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    title = doc.add_paragraph(payload['title'], 'Title')
    title.alignment = 1
    for run in title.runs:
        run.font.size = Pt(18)
    for line in payload['content'].splitlines():
        text = line.strip()
        if not text:
            continue
        text = re.sub(r'^#{1,6}\s+', '', text)
        if re.fullmatch(r'[一二三四五六七八九十]+[、．][^。；！？\n]{1,36}', text):
            p = doc.add_paragraph(text, 'Heading 1')
            p.paragraph_format.keep_with_next = True
        else:
            p = doc.add_paragraph(text)
            p.paragraph_format.first_line_indent = Pt(24)
    doc.add_page_break()
    doc.add_heading('核验与来源记录 工作附件', 1)
    doc.add_paragraph('本附件记录生成及核验边界，不作为对制度效力、事实真实性或业务审批的自动确认。')
    if payload.get('stale'):
        doc.add_paragraph('注意：事项依据或口径已变化，本稿需重新复核。')
    if payload.get('evidenceCheck') == 'edited_requires_review':
        doc.add_paragraph('正文经修改，以下为生成时依据，修改后的对应关系尚需复核。')
    doc.add_paragraph(payload.get('boundary', ''))
    if payload.get('missing'):
        doc.add_heading('待补及待确认事项', 2)
        for i, item in enumerate(payload['missing'], 1):
            doc.add_paragraph(f'{i}. {item}')
    doc.add_heading('选定材料', 2)
    source_numbers = {}
    for i, source in enumerate(payload.get('sourceSnapshot', []), 1):
        source_numbers[source['id']] = i
        doc.add_paragraph(f"材料{i} {source['title']}；版本：{source.get('version', '待说明')}；{source.get('scope', '')}")
    doc.add_heading('原文依据', 2)
    seen = set()
    for paragraph in payload.get('paragraphs', []):
        for ev in paragraph.get('evidence', []):
            key = (ev['sourceId'], ev['unitId'], ev['quote'])
            if key in seen:
                continue
            seen.add(key)
            location = doc.add_paragraph(f"材料{source_numbers.get(ev['sourceId'], '')} {ev.get('location', '')}")
            location.paragraph_format.keep_with_next = True
            location.paragraph_format.space_after = Pt(2)
            for run in location.runs:
                run.bold = True
                run.font.size = Pt(10)
            quote = doc.add_paragraph(re.sub(r'\*\*(.+?)\*\*', r'\1', ev['quote']))
            quote.paragraph_format.line_spacing = 1.25
            for run in quote.runs:
                run.font.size = Pt(11)
    doc.save(destination)
    check = Document(destination)
    assert check.paragraphs[0].text == payload['title']
    for line in payload['content'].splitlines():
        text = re.sub(r'^#{1,6}\s+', '', line.strip())
        if text:
            assert any(p.text == text for p in check.paragraphs), f'Missing paragraph: {text[:40]}'
    Path(destination).with_suffix('.validation.json').write_text(json.dumps({
        'structure': 'passed', 'paragraphs': len(check.paragraphs), 'visual': 'not_run',
        'contentCheck': payload.get('contentCheck', 'requires_review'), 'outputId': payload['id']
    }, ensure_ascii=False, indent=2), encoding='utf-8')


if __name__ == '__main__':
    build(json.loads(Path(sys.argv[1]).read_text(encoding='utf-8-sig')), sys.argv[2])
