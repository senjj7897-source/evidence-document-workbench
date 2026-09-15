from __future__ import annotations

import json
import io
import os
import re
import sys
import zipfile
from collections import Counter
from datetime import UTC, date, datetime
from pathlib import Path
from xml.etree import ElementTree as ET


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
CP = "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}"
DC = "{http://purl.org/dc/elements/1.1/}"
DCTERMS = "{http://purl.org/dc/terms/}"


def serializable(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def node_text(node):
    parts = []
    for item in node.iter():
        if item.tag in (f"{W}t", f"{W}delText", f"{W}tab"):
            parts.append("\t" if item.tag == f"{W}tab" else (item.text or ""))
        elif item.tag == f"{W}br":
            parts.append("\n")
    return "".join(parts).strip()


def attr(node, name, default=None):
    return node.attrib.get(f"{W}{name}", default) if node is not None else default


def paragraph_format(paragraph):
    ppr = paragraph.find(f"{W}pPr")
    style = attr(ppr.find(f"{W}pStyle") if ppr is not None else None, "val", "Normal")
    numbering = None
    if ppr is not None:
        num_pr = ppr.find(f"{W}numPr")
        if num_pr is not None:
            numbering = {
                "level": attr(num_pr.find(f"{W}ilvl"), "val"),
                "numId": attr(num_pr.find(f"{W}numId"), "val"),
            }
    signatures = []
    for run in paragraph.findall(f".//{W}r"):
        rpr = run.find(f"{W}rPr")
        if rpr is None:
            signatures.append("default")
            continue
        fonts = rpr.find(f"{W}rFonts")
        size = rpr.find(f"{W}sz")
        signatures.append("|".join([
            attr(fonts, "eastAsia", attr(fonts, "ascii", "default")),
            attr(size, "val", "default"),
            "b" if rpr.find(f"{W}b") is not None else "",
            "i" if rpr.find(f"{W}i") is not None else "",
        ]))
    return {
        "style": style,
        "numbering": numbering,
        "runFormats": dict(Counter(signatures)),
    }


def extract_docx(path: Path):
    result = {
        "kind": "docx",
        "blocks": [],
        "comments": [],
        "revisions": {"insertions": [], "deletions": [], "trackRevisions": False},
        "embeddedFiles": [],
        "embeddedWorkbooks": [],
        "metadata": {},
    }
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        result["embeddedFiles"] = sorted(name for name in names if name.startswith("word/embeddings/"))
        for embedded_name in result["embeddedFiles"]:
            if not embedded_name.lower().endswith((".xlsx", ".xlsm")):
                continue
            try:
                import openpyxl
                payload = archive.read(embedded_name)
                workbook = openpyxl.load_workbook(io.BytesIO(payload), data_only=False, read_only=False)
                cached_workbook = openpyxl.load_workbook(io.BytesIO(payload), data_only=True, read_only=False)
                embedded = workbook_payload(workbook, cached_workbook)
                embedded.update({
                    "name": Path(embedded_name).name,
                    "packagePath": embedded_name,
                    "size": len(payload),
                })
                result["embeddedWorkbooks"].append(embedded)
            except Exception as error:
                result["embeddedWorkbooks"].append({
                    "name": Path(embedded_name).name,
                    "packagePath": embedded_name,
                    "error": str(error),
                    "sheets": [],
                })
        if "docProps/core.xml" in names:
            core = ET.fromstring(archive.read("docProps/core.xml"))
            for key, tag in {
                "title": f"{DC}title",
                "subject": f"{DC}subject",
                "creator": f"{DC}creator",
                "lastModifiedBy": f"{CP}lastModifiedBy",
                "created": f"{DCTERMS}created",
                "modified": f"{DCTERMS}modified",
            }.items():
                element = core.find(tag)
                if element is not None and element.text:
                    result["metadata"][key] = element.text

        comments_by_id = {}
        if "word/comments.xml" in names:
            comments_root = ET.fromstring(archive.read("word/comments.xml"))
            for comment in comments_root.findall(f"{W}comment"):
                comment_id = attr(comment, "id")
                payload = {
                    "id": comment_id,
                    "author": attr(comment, "author"),
                    "date": attr(comment, "date"),
                    "text": node_text(comment),
                    "anchors": [],
                }
                comments_by_id[comment_id] = payload
                result["comments"].append(payload)

        if "word/settings.xml" in names:
            settings = ET.fromstring(archive.read("word/settings.xml"))
            result["revisions"]["trackRevisions"] = settings.find(f".//{W}trackRevisions") is not None

        document = ET.fromstring(archive.read("word/document.xml"))
        body = document.find(f"{W}body")
        page_hint = 1
        paragraph_index = 0
        table_index = 0
        block_index = 0
        for child in list(body) if body is not None else []:
            if child.tag == f"{W}p":
                paragraph_index += 1
                block_index += 1
                text = node_text(child)
                fmt = paragraph_format(child)
                comment_ids = sorted({attr(n, "id") for n in child.iter() if n.tag in (
                    f"{W}commentRangeStart", f"{W}commentReference"
                ) and attr(n, "id") is not None})
                location = {
                    "kind": "paragraph",
                    "paragraph": paragraph_index,
                    "pageHint": page_hint,
                    "label": f"第{paragraph_index}段",
                }
                if text or comment_ids:
                    block = {
                        "id": f"p-{paragraph_index}",
                        "index": block_index,
                        "type": "paragraph",
                        "text": text,
                        "location": location,
                        "format": fmt,
                        "commentIds": comment_ids,
                    }
                    result["blocks"].append(block)
                    for comment_id in comment_ids:
                        if comment_id in comments_by_id:
                            comments_by_id[comment_id]["anchors"].append(location)
                for inserted in child.findall(f".//{W}ins"):
                    inserted_text = node_text(inserted)
                    if inserted_text:
                        result["revisions"]["insertions"].append({"text": inserted_text, "location": location})
                for deleted in child.findall(f".//{W}del"):
                    deleted_text = node_text(deleted)
                    if deleted_text:
                        result["revisions"]["deletions"].append({"text": deleted_text, "location": location})
                explicit_breaks = sum(1 for node in child.iter() if (
                    node.tag == f"{W}lastRenderedPageBreak" or
                    (node.tag == f"{W}br" and attr(node, "type") == "page")
                ))
                page_hint += explicit_breaks
            elif child.tag == f"{W}tbl":
                table_index += 1
                block_index += 1
                rows = []
                for row_index, row in enumerate(child.findall(f"{W}tr"), start=1):
                    cells = []
                    for column_index, cell in enumerate(row.findall(f"{W}tc"), start=1):
                        cells.append({
                            "row": row_index,
                            "column": column_index,
                            "address": f"R{row_index}C{column_index}",
                            "text": node_text(cell),
                        })
                    rows.append(cells)
                preview = " | ".join(cell["text"] for row in rows[:3] for cell in row[:6] if cell["text"])
                result["blocks"].append({
                    "id": f"table-{table_index}",
                    "index": block_index,
                    "type": "table",
                    "text": preview,
                    "rows": rows,
                    "location": {
                        "kind": "table",
                        "table": table_index,
                        "pageHint": page_hint,
                        "label": f"表格{table_index}",
                    },
                })
        result["statistics"] = {
            "paragraphs": paragraph_index,
            "tables": table_index,
            "comments": len(result["comments"]),
            "insertions": len(result["revisions"]["insertions"]),
            "deletions": len(result["revisions"]["deletions"]),
            "embeddedFiles": len(result["embeddedFiles"]),
            "embeddedWorkbooks": len(result["embeddedWorkbooks"]),
            "embeddedWorkbookCells": sum(
                workbook.get("statistics", {}).get("nonEmptyCells", 0)
                for workbook in result["embeddedWorkbooks"]
            ),
        }
    return result


def workbook_payload(workbook, cached_workbook):
    sheets = []
    total_cells = 0
    for sheet in workbook.worksheets:
        cached_sheet = cached_workbook[sheet.title]
        cells = []
        formula_count = 0
        truncated = False
        max_nonempty = 30000
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                if len(cells) >= max_nonempty:
                    truncated = True
                    break
                is_formula = cell.data_type == "f" or (isinstance(value, str) and value.startswith("="))
                cached_value = cached_sheet[cell.coordinate].value if is_formula else value
                formula_count += int(is_formula)
                cells.append({
                    "address": cell.coordinate,
                    "row": cell.row,
                    "column": cell.column,
                    "value": serializable(cached_value if is_formula and cached_value is not None else value),
                    "formula": str(value) if is_formula else None,
                    "cachedValue": serializable(cached_value),
                    "numberFormat": cell.number_format,
                    "dataType": cell.data_type,
                    "rowHidden": bool(sheet.row_dimensions[cell.row].hidden),
                    "columnHidden": bool(sheet.column_dimensions[cell.column_letter].hidden),
                })
            if truncated:
                break
        total_cells += len(cells)
        sheets.append({
            "name": sheet.title,
            "state": sheet.sheet_state,
            "maxRow": sheet.max_row,
            "maxColumn": sheet.max_column,
            "formulaCount": formula_count,
            "cells": cells,
            "mergedRanges": [str(item) for item in sheet.merged_cells.ranges],
            "hiddenRows": [index for index, item in sheet.row_dimensions.items() if item.hidden],
            "hiddenColumns": [index for index, item in sheet.column_dimensions.items() if item.hidden],
            "truncated": truncated,
        })
    defined_names = []
    try:
        for item in workbook.defined_names.values():
            defined_names.append({"name": item.name, "value": item.attr_text})
    except Exception:
        pass
    result = {
        "kind": "xlsx",
        "sheets": sheets,
        "definedNames": defined_names,
        "statistics": {
            "sheets": len(sheets),
            "hiddenSheets": sum(sheet["state"] != "visible" for sheet in sheets),
            "nonEmptyCells": total_cells,
            "formulas": sum(sheet["formulaCount"] for sheet in sheets),
        },
    }
    workbook.close()
    cached_workbook.close()
    return result


def extract_xlsx(path: Path):
    try:
        import openpyxl
    except Exception as error:
        return {"kind": "xlsx", "error": f"openpyxl unavailable: {error}", "sheets": []}

    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    cached_workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
    return workbook_payload(workbook, cached_workbook)


def extract_pdf(path: Path):
    try:
        from pypdf import PdfReader
        reader = PdfReader(path)
        pages = []
        for index, page in enumerate(reader.pages, start=1):
            pages.append({
                "page": index,
                "text": (page.extract_text() or "").strip(),
                "location": {"kind": "page", "page": index, "label": f"第{index}页"},
            })
        return {
            "kind": "pdf",
            "pages": pages,
            "statistics": {"pages": len(pages), "textPages": sum(bool(p["text"]) for p in pages)},
            "metadata": {key.lstrip("/"): serializable(value) for key, value in (reader.metadata or {}).items()},
        }
    except Exception as error:
        return {"kind": "pdf", "error": str(error), "pages": [], "statistics": {"pages": 0, "textPages": 0}}


def extract_text(path: Path):
    text = path.read_text(encoding="utf-8", errors="replace")
    return {
        "kind": "text",
        "blocks": [
            {
                "id": f"line-{index}",
                "type": "line",
                "text": line,
                "location": {"kind": "line", "line": index, "label": f"第{index}行"},
            }
            for index, line in enumerate(text.splitlines(), start=1) if line.strip()
        ],
        "statistics": {"lines": len(text.splitlines())},
    }


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: extract.py INPUT OUTPUT")
    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    extension = source.suffix.lower()
    if extension in (".docx", ".docm"):
        payload = extract_docx(source)
    elif extension in (".xlsx", ".xlsm"):
        payload = extract_xlsx(source)
    elif extension == ".pdf":
        payload = extract_pdf(source)
    else:
        payload = extract_text(source)
    payload.update({
        "sourceName": source.name,
        "sourceSize": source.stat().st_size,
        "extractedAt": datetime.now(UTC).isoformat(timespec="seconds").replace("+00:00", "Z"),
    })
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"ok": True, "kind": payload.get("kind"), "statistics": payload.get("statistics", {})}, ensure_ascii=False))


if __name__ == "__main__":
    main()
